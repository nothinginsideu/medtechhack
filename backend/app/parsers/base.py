import abc
import typing
from pydantic import BaseModel, Field, field_validator
from decimal import Decimal

class ParsedItem(BaseModel):
    service_name_raw: str
    price_resident_kzt: Decimal | None = None
    price_nonresident_kzt: Decimal | None = None
    currency_original: str = "KZT"
    service_code_source: str | None = None
    
    @field_validator('price_resident_kzt')
    def check_price(cls, v):
        if v is not None and v <= Decimal('0'):
            raise ValueError("Цена должна быть больше 0")
        return v
        
    def validate_prices(self):
        # Если есть обе цены, проверим что нерезидент >= резидент
        if self.price_resident_kzt and self.price_nonresident_kzt:
            if self.price_nonresident_kzt < self.price_resident_kzt:
                raise ValueError("Цена нерезидента не может быть меньше цены резидента")


def detect_currency(text: str = "", filename: str = "") -> str:
    """
    Detects the currency (USD, RUB, or defaults to KZT) from a string and/or filename.
    """
    t_lower = text.lower() if text else ""
    f_lower = filename.lower() if filename else ""
    
    if "$" in t_lower or "usd" in t_lower or "$" in f_lower or "usd" in f_lower or "dollar" in f_lower or "доллар" in f_lower:
        return "USD"
    if "rub" in t_lower or "руб" in t_lower or "rub" in f_lower or "руб" in f_lower or "рубл" in f_lower:
        return "RUB"
    return "KZT"


class BaseParser(abc.ABC):
    """
    Базовый класс для всех парсеров прайс-листов.
    """
    def __init__(self, file_path: str, config: dict):
        self.file_path = file_path
        self.config = config

    @abc.abstractmethod
    def parse(self) -> typing.List[ParsedItem]:
        """
        Парсит файл и возвращает список найденных услуг с ценами.
        """
        pass


def extract_date_from_file(file_path: str) -> typing.Any:
    import os
    from datetime import date
    import re
    ext = os.path.splitext(file_path)[1].lower()
    
    # helper regex search
    def search_text_for_date(text: str) -> date | None:
        if not text:
            return None
        # 1. Look for explicit "действует с/от DD.MM.YYYY"
        match = re.search(r'(?:действует\s+с|от|утвержден|дата|с)\s*[:.-]?\s*(\d{1,2})[./-](\d{1,2})[./-](20\d{2})', text, re.IGNORECASE)
        if match:
            d, m, y = map(int, match.groups())
            try:
                return date(y, m, d)
            except ValueError:
                pass
        # 2. Look for any DD.MM.YYYY
        match = re.search(r'\b(\d{1,2})[./-](\d{1,2})[./-](20\d{2})\b', text)
        if match:
            d, m, y = map(int, match.groups())
            try:
                return date(y, m, d)
            except ValueError:
                pass
        return None

    try:
        if ext in [".xlsx", ".xls"]:
            # Excel - scan up to 30 rows
            if ext == ".xls":
                import xlrd
                wb = xlrd.open_workbook(file_path)
                sheet = wb.sheet_by_index(0)
                for i in range(min(sheet.nrows, 30)):
                    for cell in sheet.row_values(i):
                        if isinstance(cell, str):
                            d = search_text_for_date(cell)
                            if d:
                                return d
            else:
                import openpyxl
                wb = openpyxl.load_workbook(file_path, data_only=True)
                sheet = wb.active
                for row in sheet.iter_rows(max_row=30, values_only=True):
                    for cell in row:
                        if isinstance(cell, str):
                            d = search_text_for_date(cell)
                            if d:
                                return d
        elif ext == ".docx":
            # Word - scan up to 30 paragraphs
            import docx
            doc = docx.Document(file_path)
            for p in doc.paragraphs[:30]:
                d = search_text_for_date(p.text)
                if d:
                    return d
        elif ext == ".pdf":
            # PDF - scan only the first 3 pages (strict limit to avoid timeout)
            import fitz
            doc = fitz.open(file_path)
            for page in doc[:3]:
                text = page.get_text()
                d = search_text_for_date(text)
                if d:
                    return d
    except Exception as e:
        print(f"Ошибка при извлечении даты из файла: {e}")
        
    return None
